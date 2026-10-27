import os
import tempfile
import io
import re
from typing import Tuple, List, Dict
from fastapi import UploadFile, HTTPException
from app.core.config import settings

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}
MAX_FILE_BYTES = settings.MAX_UPLOAD_MB * 1024 * 1024

class DocumentParserService:
    @staticmethod
    def validate_file(file: UploadFile) -> Tuple[str, str]:
        filename = file.filename or "unknown"
        _, ext = os.path.splitext(filename.lower())
        
        if ext not in ALLOWED_EXTENSIONS:
            raise HTTPException(
                status_code=400,
                detail=f"Unsupported file format '{ext}'. Allowed formats are: PDF (.pdf), Word (.docx), Plain Text (.txt)"
            )
        return filename, ext

    @classmethod
    async def extract_text(cls, file: UploadFile) -> Dict[str, any]:
        filename, ext = cls.validate_file(file)
        
        # Read content with size check
        content = await file.read()
        if len(content) > MAX_FILE_BYTES:
            raise HTTPException(
                status_code=413,
                detail=f"File exceeds maximum allowed size of {settings.MAX_UPLOAD_MB}MB."
            )
        if len(content) == 0:
            raise HTTPException(
                status_code=400,
                detail="Uploaded file is empty (0 bytes). Please upload a valid document."
            )

        extracted_text = ""
        pages_info: List[Dict[str, any]] = []

        # Temporary directory for secure processing with guaranteed cleanup
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = os.path.join(temp_dir, f"temp_upload{ext}")
            with open(temp_path, "wb") as f:
                f.write(content)

            try:
                if ext == ".txt":
                    try:
                        extracted_text = content.decode("utf-8")
                    except UnicodeDecodeError:
                        extracted_text = content.decode("latin-1", errors="replace")
                    pages_info.append({"page_number": 1, "text": extracted_text[:1000]})

                elif ext == ".docx":
                    import docx
                    doc = docx.Document(temp_path)
                    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
                    # Also extract text from tables
                    for table in doc.tables:
                        for row in table.rows:
                            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                            if row_text:
                                paragraphs.append(row_text)
                    extracted_text = "\n\n".join(paragraphs)
                    pages_info.append({"page_number": 1, "text": extracted_text[:1000]})

                elif ext == ".pdf":
                    try:
                        import fitz  # PyMuPDF
                        pdf_doc = fitz.open(temp_path)
                        for page_num in range(len(pdf_doc)):
                            page = pdf_doc[page_num]
                            p_text = page.get_text()
                            if p_text.strip():
                                pages_info.append({
                                    "page_number": page_num + 1,
                                    "text": p_text
                                })
                                extracted_text += f"\n--- Page {page_num + 1} ---\n" + p_text
                        pdf_doc.close()
                    except ImportError:
                        # Fallback simple parser if PyMuPDF still installing
                        extracted_text = re.sub(r'[^\x20-\x7E\n\r\t]', ' ', content.decode('latin-1', errors='ignore'))
                        pages_info.append({"page_number": 1, "text": extracted_text[:1000]})

            except Exception as e:
                raise HTTPException(
                    status_code=422,
                    detail=f"Could not read content from document '{filename}'. Ensure the file is not password-protected or corrupted."
                )

        cleaned_text = re.sub(r'\s+', ' ', extracted_text).strip()
        if len(cleaned_text) < 15:
            raise HTTPException(
                status_code=422,
                detail="Extracted text is empty or too short. If this is a scanned document without an embedded text layer, OCR is required."
            )

        return {
            "filename": filename,
            "raw_text": extracted_text,
            "cleaned_text": cleaned_text,
            "char_count": len(extracted_text),
            "pages_count": max(len(pages_info), 1),
            "pages_info": pages_info
        }
